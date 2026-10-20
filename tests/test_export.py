"""Выгрузка протокола в DOCX и PDF: один разбор в блоки, два файла из одних блоков."""
import io

import docx
from docx.oxml.ns import qn
from pypdf import PdfReader

from app import export
from app.export import Code, Heading, ListBlock, Para, Quote, Rule, TableBlock

BASE = "https://podyachiy.example"
RECORD = {"file": "Проекты/Склад/планёрка 0610.mp4", "title": "Планёрка по складу: итоги/план",
          "record_date": "2026-10-06", "finished_at": "2026-10-07T09:00:00+03:00", "audio_duration_sec": 5246.6}

MD = """<!-- протокол: модель x, шаблон meeting -->

# Протокол совещания

**Название:** Планёрка по складу
**Дата:** 06.10.2026 (из имени файла)

## 1. Кратко

Обсудили *переезд* склада и `ОбменДанными`: срок **не меняется** [00:01:51], стрелки → ↔ на месте.

- Каталог: карточки заполняются. [00:02:37]
- Склад:
  - стеллажи заказаны;
  - **приёмка** — с понедельника
    и до конца месяца.
- Итог.

> Примечание: метки S2 и S3, вероятно, один человек.

### Порядок работ

1. Собрать заявки.
2. Сверить остатки.

Отдельно:

7. Закрыть период.

**Снабжение — Орлова** [00:10:39]
- Заявок стало меньше.

---

## 5. Задачи

| Задача | Срок | В записи | Статус |
|---|---|---|---|
| **Иванов Иван** | | | |
| Согласовать смету по складу | до пятницы | [00:31:23] | — |
| Показать `ТМЦ` | сегодня | [01:02:03] | в работе |
| **Ответственный не назван** | | | |
| Опубликовать ссылку | — | [00:17:43] |

```
Запрос.УстановитьПараметр("Склад", Склад);
  Запрос.Выполнить();
```

## Поручения прошлого совещания

Незакрытые поручения из протокола записи [«Планёрка 29.09»](/?path=%D0%9F&open=a.mp4). Раздел добавлен сервисом.

- **Петров:** сдать отчёт; [в записи 00:05:00](/?path=%D0%9F&open=a.mp4&t=300); см. [регламент](https://example.org/doc)
"""


def kinds(blocks):
    return [type(b).__name__ for b in blocks]


def test_protocol_is_parsed_into_blocks_of_every_kind():
    blocks = export.parse(MD)
    assert kinds(blocks) == ["Heading", "Para", "Heading", "Para", "ListBlock", "Quote", "Heading", "ListBlock",
                             "Para", "ListBlock", "Para", "ListBlock", "Rule", "Heading", "TableBlock", "Code",
                             "Heading", "Para", "ListBlock"]
    assert [(b.level, export.text_of(b.spans)) for b in blocks if isinstance(b, Heading)] == [
        (1, "Протокол совещания"), (2, "1. Кратко"), (3, "Порядок работ"), (2, "5. Задачи"),
        (2, "Поручения прошлого совещания")]
    assert "<!--" not in "".join(export.text_of(line) for b in blocks if isinstance(b, Para) for line in b.lines)

    # строки сведений остаются отдельными строками одного абзаца
    facts = blocks[1]
    assert [export.text_of(line) for line in facts.lines] == ["Название: Планёрка по складу",
                                                              "Дата: 06.10.2026 (из имени файла)"]
    assert facts.lines[0][0].bold and not facts.lines[0][1].bold

    # жирный, курсив, код и метка времени — отдельными кусками
    spans = blocks[3].lines[0]
    marked = {s.text: (s.bold, s.italic, s.code, s.clock) for s in spans}
    assert marked["переезд"] == (False, True, False, False)
    assert marked["ОбменДанными"] == (False, False, True, False)
    assert marked["не меняется"] == (True, False, False, False)
    assert marked["[00:01:51]"] == (False, False, False, True)
    assert export.text_of(spans).endswith("стрелки → ↔ на месте.")

    # список: вложенность, продолжение пункта на следующей строке
    items = blocks[4].items
    assert [(i.level, i.ordered) for i in items] == [(0, False), (0, False), (1, False), (1, False), (0, False)]
    assert export.text_of(items[3].spans) == "приёмка — с понедельника и до конца месяца."
    assert items[3].spans[0].bold
    # нумерованные списки помнят свои номера: второй начинается с 7, а не с 1
    assert [[i.number for i in b.items] for b in blocks if isinstance(b, ListBlock) and b.items[0].ordered] == [[1, 2], [7]]

    assert [export.text_of(line) for line in blocks[5].lines] == ["Примечание: метки S2 и S3, вероятно, один человек."]
    assert export.is_lead(blocks[10]) and not export.is_lead(blocks[8])       # «**Тема — докладчик** [время]»
    assert isinstance(blocks[12], Rule)
    assert blocks[15].text == 'Запрос.УстановитьПараметр("Склад", Склад);\n  Запрос.Выполнить();'


def test_table_keeps_the_owner_rows_and_gets_even_rows():
    table = next(b for b in export.parse(MD) if isinstance(b, TableBlock))
    assert [export.text_of(c) for c in table.header] == ["Задача", "Срок", "В записи", "Статус"]
    assert [r.group for r in table.rows] == [True, False, False, True, False]
    assert all(len(r.cells) == 4 for r in table.rows)                    # короткая строка дополнена пустой ячейкой
    assert export.text_of(table.rows[0].cells[0]) == "Иванов Иван"
    assert table.rows[2].cells[0][1].code and table.rows[1].cells[2][0].clock
    shares = export.shares(table)
    assert abs(sum(shares) - 1) < 1e-9 and shares[0] == max(shares) and min(shares) > 0.1


def test_links_need_the_permanent_address_of_the_service():
    """Без постоянного адреса метки времени и ссылки на другие записи остаются текстом;
    внешняя ссылка работает всегда."""
    def hrefs(base):
        blocks = export.parse(MD, base, RECORD["file"])
        return [s.href for b in blocks for line in (b.lines if isinstance(b, Para) else [i.spans for i in b.items]
                                                    if isinstance(b, ListBlock) else []) for s in line if s.href]

    assert hrefs("") == ["https://example.org/doc"]
    linked = hrefs(BASE)
    assert linked[0] == (BASE + "/?path=%D0%9F%D1%80%D0%BE%D0%B5%D0%BA%D1%82%D1%8B%2F%D0%A1%D0%BA%D0%BB%D0%B0%D0%B4"
                         "&open=%D0%9F%D1%80%D0%BE%D0%B5%D0%BA%D1%82%D1%8B%2F%D0%A1%D0%BA%D0%BB%D0%B0%D0%B4%2F"
                         "%D0%BF%D0%BB%D0%B0%D0%BD%D1%91%D1%80%D0%BA%D0%B0%200610.mp4&t=111")
    assert linked[-3:] == [BASE + "/?path=%D0%9F&open=a.mp4", BASE + "/?path=%D0%9F&open=a.mp4&t=300",
                           "https://example.org/doc"]


def test_document_takes_the_title_and_facts_from_the_record():
    doc = export.document(MD, RECORD)
    assert doc.title == "Планёрка по складу: итоги/план"
    assert doc.facts == [("Дата совещания", "06.10.2026"), ("Длительность", "1 ч 27 мин"),
                         ("Папка", "Проекты/Склад"), ("Запись", "планёрка 0610.mp4")]
    assert kinds(doc.blocks)[0] == "Heading"                              # заголовок протокола другой — остаётся

    # названия в журнале нет — берётся заголовок протокола и из текста убирается, чтобы не стоять дважды
    bare = export.document(MD, {"file": "запись.mp4", "finished_at": "2026-10-07T09:00:00+03:00"})
    assert bare.title == "Протокол совещания" and kinds(bare.blocks)[0] == "Para"
    assert bare.facts == [("Обработано", "07.10.2026"), ("Запись", "запись.mp4")]
    assert export.document("Текст.\n", {"file": "Папка/запись 1.mp4"}).title == "запись 1"
    assert export.clock(1911) == "32 мин" and export.clock(20) == "1 мин" and export.clock(3600) == "1 ч 00 мин"


def test_file_name_is_the_title_without_forbidden_signs():
    assert export.file_name("Планёрка по складу: итоги/план", "docx") == "Планёрка по складу итоги план.docx"
    assert export.file_name('  «Отчёт» <1>?. ', "pdf") == "«Отчёт» 1.pdf"
    assert export.file_name("///", "pdf") == "протокол.pdf"
    assert len(export.file_name("я" * 300, "pdf")) == 124


def test_docx_has_word_headings_lists_and_a_real_table():
    data, name = export.build(MD, RECORD, "docx", BASE)
    assert name == "Планёрка по складу итоги план.docx"
    doc = docx.Document(io.BytesIO(data))
    pars = [(p.style.name, p.text) for p in doc.paragraphs]
    assert pars[0] == ("Title", "Планёрка по складу: итоги/план")
    assert pars[1] == ("Protocol Facts", "Дата совещания: 06.10.2026 · Длительность: 1 ч 27 мин · "
                                         "Папка: Проекты/Склад · Запись: планёрка 0610.mp4")
    assert [p for p in pars if p[0].startswith("Heading")] == [
        ("Heading 1", "Протокол совещания"), ("Heading 2", "1. Кратко"), ("Heading 3", "Порядок работ"),
        ("Heading 2", "5. Задачи"), ("Heading 2", "Поручения прошлого совещания")]
    assert [p[0] for p in pars if p[0].startswith("List")] == [
        "List Bullet", "List Bullet", "List Bullet 2", "List Bullet 2", "List Bullet", "List Number", "List Number",
        "List Number", "List Bullet", "List Bullet"]
    assert ("Quote", "Примечание: метки S2 и S3, вероятно, один человек.") in pars
    assert ("Protocol Code", "  Запрос.Выполнить();") in pars
    assert doc.core_properties.title == "Планёрка по складу: итоги/план" and doc.core_properties.author == "Подьячий"

    section = doc.sections[0]
    assert (round(section.page_width.cm, 1), round(section.page_height.cm, 1)) == (21.0, 29.7)       # А4
    assert "PAGE" in section.footer._element.xml and "NUMPAGES" in section.footer._element.xml

    # нумерованные списки не продолжают друг друга: у каждого свой счётчик со своим началом
    starts = {}
    for num in doc.part.numbering_part.element.findall(qn("w:num")):
        for over in num.iter(qn("w:startOverride")):
            starts[num.numId] = over.val
    numbered = [p._p.pPr.numPr.numId.val for p in doc.paragraphs if p.style.name == "List Number"]
    assert numbered[0] == numbered[1] != numbered[2]
    assert [starts[n] for n in numbered] == [1, 1, 7]

    table, = doc.tables
    assert len(table.rows) == 6 and len(table.columns) == 4
    assert [c.text for c in table.rows[0].cells] == ["Задача", "Срок", "В записи", "Статус"]
    assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None      # шапка повторяется на страницах
    assert len({id(c._tc) for c in table.rows[1].cells}) == 1 and table.rows[1].cells[0].text == "Иванов Иван"
    assert [c.text for c in table.rows[2].cells] == ["Согласовать смету по складу", "до пятницы", "[00:31:23]", "—"]
    assert abs(sum(c.width for c in table.rows[2].cells) - section.page_width
               + section.left_margin + section.right_margin) < 2000     # таблица ровно в ширину полосы

    links = [r.target_ref for r in doc.part.rels.values() if r.reltype.endswith("/hyperlink")]
    assert "https://example.org/doc" in links and any(link.endswith("&t=1883") for link in links)
    plain = docx.Document(io.BytesIO(export.build(MD, RECORD, "docx")[0]))
    assert [r.target_ref for r in plain.part.rels.values() if r.reltype.endswith("/hyperlink")] == ["https://example.org/doc"]
    assert "[00:31:23]" in plain.tables[0].rows[2].cells[2].text


def test_pdf_is_a_real_file_with_cyrillic_text_title_and_outline():
    data, name = export.build(MD, RECORD, "pdf", BASE)
    assert name.endswith(".pdf") and data.startswith(b"%PDF-") and data.rstrip().endswith(b"%%EOF")
    pdf = PdfReader(io.BytesIO(data))
    assert len(pdf.pages) == 2 or len(pdf.pages) == 1
    assert pdf.metadata.title == "Планёрка по складу: итоги/план" and pdf.metadata.author == "Подьячий"
    text = "\n".join(page.extract_text() for page in pdf.pages)
    for piece in ("Планёрка по складу: итоги/план", "Дата совещания: 06.10.2026", "5. Задачи", "Иванов Иван",
                  "Согласовать смету по складу", "[00:31:23]", "стрелки → ↔", "7. Закрыть период",
                  "Запрос.Выполнить();", f"Страница 1 из {len(pdf.pages)}"):
        assert piece in text, piece
    assert [o.title for o in pdf.outline if not isinstance(o, list)] == ["Планёрка по складу: итоги/план"]
    fonts = {f["/BaseFont"] for page in pdf.pages for f in (r.get_object() for r in page["/Resources"]["/Font"].values())}
    assert any("DejaVuSans" in f for f in fonts)                          # шрифт с кириллицей встроен
    links = [a.get_object()["/A"]["/URI"] for page in pdf.pages for a in page.get("/Annots", [])]
    assert "https://example.org/doc" in links and any(link.endswith("&t=1883") for link in links)


def test_long_table_goes_to_the_next_pages_with_its_header():
    rows = "".join(f"| Задача номер {n}: {'проверить остатки и сверить заявки; ' * 6} | до {n}.11 | [00:{n:02d}:00] | — |\n"
                   for n in range(1, 60))
    md = "## Задачи\n\n| Задача | Срок | В записи | Статус |\n|---|---|---|---|\n" + rows
    pdf = PdfReader(io.BytesIO(export.build(md, {"file": "запись.mp4"}, "pdf")[0]))
    texts = [page.extract_text() for page in pdf.pages]
    assert len(pdf.pages) >= 3 and "Задача номер 1:" in texts[0]          # таблица начинается сразу под заголовком
    assert all("Задача" in t and "В записи" in t and "Статус" in t for t in texts)        # шапка на каждой странице
    assert sum(t.count("Задача номер") for t in texts) == 59                                # строки не потеряны и не удвоены
    assert all(t.strip() for t in texts)

    # строка выше страницы делится между страницами, а не роняет сборку
    tall = "| Задача | Срок |\n|---|---|\n| " + "слово " * 3000 + " | — |\n| после | завтра |\n"
    pdf = PdfReader(io.BytesIO(export.build(tall, {"file": "запись.mp4"}, "pdf")[0]))
    text = "".join(page.extract_text() for page in pdf.pages)
    assert len(pdf.pages) >= 3 and text.count("слово") == 3000 and "после" in text
