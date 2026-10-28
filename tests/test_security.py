"""Защита сервиса в браузере и на сервере: чужие страницы, заголовки, файлы для Excel, ffmpeg.

Страница сервиса открыта в браузере рядом с чужими сайтами. Браузер сам приложит к запросу куку
входа, с какой бы страницы запрос ни ушёл, — поэтому изменяющий запрос с кукой и с чужой страницы
сервис не выполняет (CSRF). Заголовки ответов не дают странице сервиса исполнять чужие сценарии
(Content-Security-Policy), открываться внутри чужой страницы, а браузеру — хранить у себя данные
из API. Строки из протоколов и имён записей в CSV не исполняются в Excel формулами.
"""
import base64
import hashlib
import re
import subprocess
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app import costbook, media, speech
from tests.test_web import ADMIN_PASSWORD, HOME, PASSWORD, ROOT, enter, keyed, recorded, running, staffed

STATIC = ROOT / "app" / "static"
PAGES = sorted(STATIC.glob("*.html"))


@pytest.fixture
def site(tmp_path):
    with running(tmp_path, extra={"PUBLIC_URL": "https://meetings.example.test"}) as (web, client):
        yield web, client


# ---------- изменяющие запросы с чужой страницы ----------

def test_cookie_request_from_another_site_is_refused(site):
    """Кука входа уходит с любой страницы — и с соседнего поддомена (SameSite его своим считает).
    Изменяющий запрос с кукой и с чужим Origin — отказ 403; без Origin решает Sec-Fetch-Site.
    Своя страница — адрес, по которому сервис открыт, или PUBLIC_URL (за прокси)."""
    web, admin, ivanov, _ = staffed(site)
    folder = f"{HOME}/ivanov/Совещания"
    for headers in ({"origin": "https://evil.example"}, {"origin": "https://sub.meetings.example.test"},
                    {"origin": "null"}, {"sec-fetch-site": "same-site"}, {"sec-fetch-site": "cross-site"}):
        refused = ivanov.post("/api/folder", json={"dir": folder}, headers=headers)
        assert refused.status_code == 403 and refused.json()["detail"] == web.server.FOREIGN_PAGE, headers
        # параметры в адресе, без тела — такой запрос чужая страница шлёт обычной формой
        moved = ivanov.post("/api/v1/speakers", params={"file": f"{HOME}/ivanov/вкс.mp4", "rebuild": "true"},
                            headers=headers)
        assert moved.status_code == 403 and moved.json()["detail"] == web.server.FOREIGN_PAGE, headers
    assert not (web.core.cfg.inbox / folder).exists()
    # чтение не меняет ничего, а ответ чужой странице браузер без CORS не отдаст
    assert ivanov.get("/api/browse", headers={"origin": "https://evil.example"}).status_code == 200
    # вход с чужой страницы тоже не принимается: иначе она входит в сервис под своей учётной записью
    assert TestClient(web.app).post("/api/login", json={"login": "ivanov", "password": PASSWORD},
                                    headers={"origin": "https://evil.example"}).status_code == 403

    for n, headers in enumerate(({"origin": "http://testserver"}, {"origin": "https://meetings.example.test"},
                                 {"origin": "https://MEETINGS.example.test:443"},
                                 {"sec-fetch-site": "same-origin"}, {"sec-fetch-site": "none"}, {})):
        made = ivanov.post("/api/folder", json={"dir": f"{folder}{n}"}, headers=headers)
        assert made.status_code == 200, (headers, made.text)


def test_requests_with_a_key_or_password_are_not_checked_for_origin(site):
    """Ключ и пароль в заголовке Authorization браузер сам не прикладывает — такой запрос
    составлен программой нарочно, и Origin в нём ничего не говорит."""
    web, admin, ivanov, _ = staffed(site)
    robot, _ = keyed(web, ivanov)
    evil = {"origin": "https://evil.example", "sec-fetch-site": "cross-site"}
    assert robot.post("/api/folder", json={"dir": f"{HOME}/ivanov/Ключом"}, headers=evil).status_code == 200
    basic = "Basic " + base64.b64encode(f"ivanov:{PASSWORD}".encode()).decode()
    assert TestClient(web.app).post("/api/folder", json={"dir": f"{HOME}/ivanov/Паролем"},
                                    headers={**evil, "authorization": basic}).status_code == 200


def test_personal_mode_checks_the_same_way(tmp_path):
    with running(tmp_path, mode="personal") as (web, client):
        for headers in ({"origin": "https://evil.example"}, {"sec-fetch-site": "same-site"}):
            refused = client.post("/api/folder", json={"dir": "Чужая"}, headers={"host": "localhost:8080", **headers})
            assert refused.status_code == 403 and refused.json()["detail"] == web.server.FOREIGN_PAGE
        assert client.post("/api/folder", json={"dir": "Своя"},
                           headers={"host": "localhost:8080", "origin": "http://localhost:8080"}).status_code == 200


# ---------- заголовки ответов ----------

def test_every_answer_carries_the_security_headers(site):
    web, admin, ivanov, _ = staffed(site)
    for path in ("/", "/login", "/static/app.css", "/api/status", "/api/browse", "/no-such-page"):
        answer = ivanov.get(path)
        assert answer.headers["content-security-policy"] == web.common.CSP, path
        assert answer.headers["x-content-type-options"] == "nosniff", path
        assert answer.headers["referrer-policy"] == "same-origin", path
        assert "strict-transport-security" not in answer.headers, path         # по HTTP HSTS не ставят
    policy = dict(part.strip().split(" ", 1) for part in web.common.CSP.split(";"))
    assert policy["script-src"] == "'self'" and policy["frame-ancestors"] == "'none'" and policy["object-src"] == "'none'"


def test_hsts_only_over_https(site):
    web, _ = site
    secure = TestClient(web.app, base_url="https://testserver")
    assert secure.get("/api/status").headers["strict-transport-security"] == "max-age=31536000"
    assert "strict-transport-security" not in TestClient(web.app).get("/api/status").headers


def test_api_answers_and_files_are_not_stored_by_the_browser(site):
    """Ответы API, выгрузки и звук браузер не кладёт в кэш на диске: на общем компьютере их
    прочёл бы следующий. Страницы и их сценарии — no-cache (сверка при открытии), логотип с
    версией в адресе — хранится год."""
    web, admin, ivanov, _ = staffed(site)
    rel = f"{HOME}/ivanov/вкс.mp4"
    recorded(web, rel)
    out = web.core.cfg.outbox / HOME / "ivanov" / "вкс"
    (out / f"вкс{media.AUDIO_SUFFIX}").write_bytes(b"ID3audio")
    for path, params in (("/api/browse", {}), ("/api/audio", {"file": rel}), ("/api/export", {"file": rel, "format": "docx"}),
                         ("/api/map", {"file": rel, "download": "true"}), ("/api/v1/tasks.csv", {}),
                         ("/api/admin/report.csv", {}), ("/api/admin/report.xlsx", {})):
        answer = (admin if "admin" in path else ivanov).get(path, params=params)
        assert answer.status_code == 200, (path, answer.text[:200])
        assert answer.headers["cache-control"] == "no-store", path
    assert ivanov.get("/static/common.js").headers["cache-control"] == "no-cache"
    assert ivanov.get("/").headers["cache-control"] == "no-cache"


def test_map_page_runs_its_own_script_and_nothing_else(site):
    """Карта — самостоятельный файл со сценарием внутри: его пересылают и открывают без сервиса.
    Сервис, показывая её, разрешает ровно этот сценарий — по хешу, а не любой встроенный."""
    web, admin, ivanov, _ = staffed(site)
    rel = f"{HOME}/ivanov/вкс.mp4"
    recorded(web, rel)
    shown = ivanov.get("/api/map", params={"file": rel})
    scripts = re.findall(r"<script>(.*?)</script>", shown.text, flags=re.S)
    assert len(scripts) == 1
    digest = base64.b64encode(hashlib.sha256(scripts[0].encode("utf-8")).digest()).decode()
    policy = shown.headers["content-security-policy"]
    assert f"script-src 'self' 'sha256-{digest}';" in policy and "'unsafe-inline'" not in policy.split("script-src")[1].split(";")[0]
    assert policy.replace(f" 'sha256-{digest}'", "") == web.common.CSP


@pytest.mark.parametrize("page", PAGES, ids=lambda p: p.name)
def test_pages_have_no_inline_scripts_or_handlers(page):
    """Сценарии страниц — отдельными файлами (CSP пускает только их). Обработчики on…= в разметке
    и в строках, которые сценарий вставляет через innerHTML, CSP тоже не исполнит — их нет."""
    text = page.read_text(encoding="utf-8")
    for tag in re.findall(r"<script\b[^>]*>", text):
        assert re.fullmatch(r'<script src="/static/[\w.-]+\.js">', tag), tag
    own = STATIC / f"{page.stem}.js"
    assert own.is_file() and f'<script src="/static/{own.name}"></script>' in text
    for source in (page, *STATIC.glob("*.js")):
        code = source.read_text(encoding="utf-8")
        assert not re.search(r"<[a-zA-Z][^<>]*\son[a-z]+\s*=", code), source.name
        assert "javascript:" not in code.lower(), source.name


def test_login_page_returns_only_inside_the_site():
    """После входа — туда, откуда отправили входить, но только на свой сайт: «/\\evil.ru», «//evil.ru»,
    «javascript:» адрес браузер увёл бы на чужой. Адрес разбирает сам браузер (new URL) и сверяет
    origin со своим; уходит он по полному адресу своего сайта, а не по пути."""
    code = (STATIC / "login.js").read_text(encoding="utf-8")
    back = re.search(r"function inside\(raw\) \{(.*?)\n\}", code, flags=re.S).group(1)
    assert "new URL(raw, location.origin)" in back and "url.origin === location.origin" in back
    assert "location.href = next" in code and "location.href = q.get" not in code


# ---------- CSV и книга Excel ----------

def test_csv_cells_never_start_a_formula():
    """Excel исполняет ячейку CSV, которая начинается с = + - @ (и с табуляции или перевода
    строки), как формулу: имя записи «=HYPERLINK(…)» стало бы ссылкой наружу. Таким строкам
    спереди ставится апостроф; числа остаются числами."""
    for raw in ("=HYPERLINK(\"http://evil.example\")", "+7 (900)", "-сделать", "@SUM(1)", "\tтаб", "\rвозврат"):
        assert costbook.cell(raw) == "'" + raw
    for raw in ("Иванов", "", "12", "1-2", "a=b"):
        assert costbook.cell(raw) == raw
    table = costbook.Table("Записи", ["Папка", "Запись", "Потрачено, $"], [["=1+1", "@x", -0.5]], {2})
    lines = costbook.to_csv([("Записи", [table])]).lstrip("﻿").split("\r\n")
    assert lines[2] == "'=1+1;'@x;-0.500000"


def test_tasks_csv_cells_never_start_a_formula(site, monkeypatch):
    web, admin, ivanov, _ = staffed(site)
    task = {"who": "=cmd|' /C calc'!A0", "absent": False, "task": "-сделать отчёт", "due": None, "due_text": "+1 день",
            "due_date": "2026-10-20", "overdue": False, "status": "@статус", "title": "=Планёрка", "date": "2026-10-06",
            "dir": "=Папка", "file": "=Папка/=вкс.mp4", "at": "00:01:02", "link": "http://testserver/", "done": False,
            "done_by": None, "done_at": None, "comment": "=1+1", "due_by": None}

    async def tasks(*args, **kwargs):
        return {"tasks": [task]}

    monkeypatch.setattr(web.api_v1, "_tasks", tasks)
    sheet = ivanov.get("/api/v1/tasks.csv").content.decode("utf-8-sig").splitlines()
    row = sheet[1].split(";")
    assert row[0] == "'=cmd|' /C calc'!A0" and row[1] == "'-сделать отчёт" and row[3] == "'+1 день"
    assert row[6] == "'@статус" and row[7] == "'=Планёрка" and row[9] == "'=Папка" and row[10] == "'=вкс.mp4"
    assert row[16] == "'=1+1" and row[8] == "2026-10-06"


def test_xlsx_keeps_formula_like_names_as_text():
    """В книге строка «=…» записана текстом (тип s): Excel её не вычисляет."""
    import io
    import openpyxl
    table = costbook.Table("Записи", ["Папка", "Запись", "Потрачено, $"], [["=1+1", "@x", 0.5]], {2})
    book = openpyxl.load_workbook(io.BytesIO(costbook.to_xlsx([("Записи", [table])])))
    cells = book["Записи"]["A2"], book["Записи"]["B2"], book["Записи"]["C2"]
    assert [c.data_type for c in cells] == ["s", "s", "n"] and cells[0].value == "=1+1"


# ---------- пути сервера в ответах ----------

def test_file_errors_do_not_show_server_paths(site, monkeypatch):
    """Сбой файловой системы — отказ 500 словами, без путей сервера: путь целиком — в журнал работы."""
    web, admin, ivanov, _ = staffed(site)
    lines = []
    monkeypatch.setattr(web.access, "log", lines.append)

    def broken(self, *args, **kwargs):
        raise PermissionError(13, "Permission denied", str(self))

    monkeypatch.setattr(web.core.Path, "mkdir", broken)
    refused = ivanov.post("/api/folder", json={"dir": f"{HOME}/ivanov/Совещания"})
    assert refused.status_code == 500
    detail = refused.json()["detail"]
    assert detail == "не удалось создать папку: нет прав на запись в папку" and str(web.core.cfg.inbox) not in detail
    assert any(str(web.core.cfg.inbox) in line for line in lines)


# ---------- ручной запуск ----------

def test_manual_run_takes_only_what_the_list_offers(site):
    """Ручной запуск берёт то, что «Обработка» показывает в списке файлов: медиафайл входной
    папки вне служебных папок (корзина .trash, недогрузки .uploads) и не временный. Файл, который
    ещё дописывается, — отказ: сервис возьмёт его сам, когда он перестанет меняться."""
    web, client = site
    run = lambda file: client.post("/api/admin/run", json={"file": file}, auth=("admin", ADMIN_PASSWORD))
    for rel in (".trash/20261001-120000/Проекты/запись.mp4", ".uploads/ab12.mp4", "Проекты/.uploads/запись.mp4",
                "Проекты/заметки.txt", "Проекты/~$запись.mp4", "Проекты/запись.mp4.part"):
        path = web.core.cfg.inbox / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(b"video")
        refused = run(rel)
        assert refused.status_code == 400 and refused.json()["detail"] == web.processing.NOT_RUNNABLE, rel
    (web.core.cfg.inbox / "Проекты" / "вкс.mp4").write_bytes(b"video")
    web.core.service.scan_once()
    assert "Проекты/вкс.mp4" in web.core.service.pending
    busy = run("Проекты/вкс.mp4")
    assert busy.status_code == 409 and busy.json()["detail"] == web.processing.STILL_WRITTEN
    assert web.core.service.queue == []
    web.core.service.pending.clear()
    assert run("Проекты/вкс.mp4").json()["queued"] == "Проекты/вкс.mp4"


# ---------- ffmpeg ----------

def calls(monkeypatch):
    seen = []

    def fake_run(args, **kwargs):
        seen.append(args)
        return subprocess.CompletedProcess(args, 0, stdout="12.5\n", stderr="max_volume: -12.0 dB")

    monkeypatch.setattr(subprocess, "run", fake_run)
    return seen


def test_ffmpeg_reads_only_local_files_and_names_the_output_format(monkeypatch, tmp_path):
    """Файл записи — не обязательно то, чем назван: плейлист (HLS, concat) под именем .mp4 заставил бы
    ffmpeg ходить в сеть или читать другие файлы. Входу разрешён только протокол file; формат
    того, что пишется, назван явно, а не угадан по имени."""
    seen = calls(monkeypatch)
    monkeypatch.setattr(media.Path, "replace", lambda self, target: None)
    src = tmp_path / "запись.mp4"
    assert media.has_audio_stream(src)
    media.ensure_audio(src, tmp_path / "out", "запись", True, lambda line: None)
    speech.make_chunk(src, tmp_path / "1.mp3", 0.0, 10.0)
    assert speech.audio_duration(src) == 12.5
    assert [args[0] for args in seen] == ["ffprobe", "ffprobe", "ffmpeg", "ffmpeg", "ffmpeg", "ffprobe"]
    for args in seen:
        assert args[1:3] == ["-protocol_whitelist", "file"], args
        if args[0] == "ffmpeg":
            out = args[-1]
            assert args[-3:-1] == ["-f", "mp3" if out.endswith(".mp3") else "null"], args
