"""Оформление корпоративного режима: логотип организации и её название.

Файл логотипа проверяется по содержимому; наружу уходят только копии, собранные сервисом.
Менять оформление может только администратор сервиса на странице; в личном режиме раздела нет.
"""
import io
import os

import pytest
from PIL import Image

from app import brand, db, mail, mindmap
from tests.test_web import (ADMIN_PASSWORD, MAIL_ENV, enter, filled, keyed, kinds, running,  # noqa: F401
                            site, staffed)


def picture(fmt="PNG", size=(512, 512), mode="RGBA", color=(20, 90, 200, 0), **save) -> bytes:
    out = io.BytesIO()
    Image.new(mode, size, color).save(out, fmt, **save)
    return out.getvalue()


SVG = b'<?xml version="1.0"?>\n<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script></svg>'


# ---------- проверка файла ----------

@pytest.mark.parametrize("fmt, mode, color", [("PNG", "RGBA", (1, 2, 3, 0)), ("JPEG", "RGB", (1, 2, 3)),
                                                ("WEBP", "RGBA", (1, 2, 3, 0))])
def test_png_jpeg_and_webp_are_accepted_by_their_content(fmt, mode, color):
    img, info = brand.read_logo(picture(fmt, (300, 300), mode, color))
    assert info["format"] == {"PNG": "PNG", "JPEG": "JPEG", "WEBP": "WebP"}[fmt]
    assert (info["width"], info["height"]) == (300, 300) and img.mode == "RGBA" and img.info == {}


@pytest.mark.parametrize("data, why", [
    (SVG, "может быть сценарий"),
    (b"<svg viewBox='0 0 1 1'></svg>", "SVG не принимается"),
    (b"\xef\xbb\xbf  <SVG></SVG>", "SVG не принимается"),
    (picture("GIF", (300, 300), "P", 0), "GIF не принимается"),
    (b"MZ\x90\x00" + b"\x00" * 600, "смотрит на содержимое файла, а не на расширение"),
    (b"%PDF-1.7 logo.png", "не картинка PNG, JPEG или WebP"),
    (b"", "файл пустой"),
    (picture(size=(255, 400)), "слишком мала: 255×400"),
    (picture(size=(4097, 256)), "слишком велика: 4097×256"),
    (picture()[:200], "повреждён или обрезан"),
    (b"\x89PNG\r\n\x1a\n" + os.urandom(64), "повреждён или обрезан"),
    (b"\x89PNG\r\n\x1a\n" + b"0" * (2 * 1024 * 1024), "больше 2 МБ"),
])
def test_everything_else_is_refused_with_the_reason(data, why):
    with pytest.raises(brand.BrandError) as refused:
        brand.read_logo(data)
    assert why in str(refused.value)


def test_borders_of_the_size_are_accepted_and_the_advice_is_not_a_refusal():
    assert brand.read_logo(picture(size=(256, 256)))[1]["advice"] == []
    wide = brand.read_logo(picture("JPEG", (4096, 256), "RGB", (255, 255, 255)))[1]["advice"]
    assert len(wide) == 2 and "не квадратный (4096×256)" in wide[0] and "непрозрачный фон" in wide[1]
    solid = brand.read_logo(picture(color=(1, 2, 3, 255)))[1]["advice"]                # PNG с альфой, но без прозрачности
    assert len(solid) == 1 and "непрозрачный фон" in solid[0]


def test_copies_are_made_by_the_service_and_carry_no_metadata(tmp_path):
    exif = Image.Exif()
    exif[0x010E] = "secret-description"
    exif[0x0112] = 6                                    # снимок с телефона: «повернуть на 90°»
    source = picture("JPEG", (600, 400), "RGB", (200, 10, 10), exif=exif.tobytes())
    assert b"secret-description" in source
    look = brand.Brand(db.Database(tmp_path))
    info = look.set_logo(source, "admin")
    assert (info["width"], info["height"]) == (400, 600)                               # поворот из метаданных применён
    version = look.view()["logo"]
    for name, box in brand.FILES.items():
        data = look.file(version, name).read_bytes()
        made = Image.open(io.BytesIO(data))
        assert made.format == "PNG" and made.width <= box[0] and made.height <= box[1], name
        assert "exif" not in made.info and b"secret-description" not in data and source[:200] not in data
    assert Image.open(look.file(version, "favicon.png")).size == (64, 64)
    # на диске — только картинки, собранные сервисом; название и сведения о логотипе — в базе
    assert sorted(p.name for p in (tmp_path / "brand").iterdir()) == ["favicon.png", "logo-small.png", "logo.png"]


def test_version_changes_with_every_logo_and_the_choice_survives_a_restart(tmp_path):
    look = brand.Brand(db.Database(tmp_path))
    assert look.view() == {"name": "", "logo": "", "by": "", "at": "", "source": {}} and look.file("", "logo.png") is None
    look.set_logo(picture(), "admin")
    first = look.view()["logo"]
    assert look.set_name("  ООО   Ромашка ", "admin") == "ООО Ромашка"
    look.set_logo(picture("WEBP"), "ivanov")
    again = brand.Brand(db.Database(tmp_path)).view()
    assert again["logo"] != first and again["name"] == "ООО Ромашка" and again["by"] == "ivanov"
    assert again["source"] == {"format": "WebP", "width": 512, "height": 512, "bytes": len(picture("WEBP"))}
    assert look.file(first, "logo.png") is None and look.file(again["logo"], "нет.png") is None
    assert look.file(again["logo"], "../brand.json") is None
    look.reset("admin")
    assert brand.Brand(db.Database(tmp_path)).view()["logo"] == "" and not (tmp_path / "brand" / "logo.png").exists()
    with pytest.raises(brand.BrandError):
        look.set_logo(SVG, "admin")
    assert look.view()["logo"] == ""


def test_lost_logo_files_fall_back_to_the_standard_logo_until_a_new_one_is_set(tmp_path):
    look = brand.Brand(db.Database(tmp_path))
    look.set_logo(picture(), "admin")
    look.set_name("ООО Ромашка", "admin")
    (tmp_path / "brand" / "logo-small.png").unlink()
    lines = []
    lost = brand.Brand(db.Database(tmp_path), lines.append)
    assert lost.view()["logo"] == "" and lost.view()["name"] == "ООО Ромашка" and "файлов логотипа" in lines[0]
    lost.set_logo(picture("WEBP"), "admin")
    assert lost.view()["logo"] and brand.Brand(db.Database(tmp_path), lines.append).view()["logo"] and len(lines) == 1


@pytest.mark.parametrize("name, why", [("я" * 41, "до 40 знаков"), ("<b>Ромашка</b>", "знаков < и >"),
                                        ("Рома\x07шка", "знаков < и >")])
def test_bad_names_are_refused(name, why):
    with pytest.raises(brand.BrandError) as refused:
        brand.clean_name(name)
    assert why in str(refused.value)
    assert brand.clean_name("я" * 40) == "я" * 40 and brand.clean_name("  ") == ""


def test_letter_carries_the_name_next_to_the_service_and_never_as_markup(monkeypatch):
    cfg = type("Cfg", (), {"smtp_from": "s@example.test"})()
    monkeypatch.setattr(mail, "organization", lambda: 'АО "Ромашка & Ко"')
    letter = mail.message(cfg, ["a@example.test"], "тема", "текст", mail._page("Заголовок", ["абзац"], None, "подвал"))
    assert letter["Subject"] == 'Подьячий · АО "Ромашка & Ко": тема'
    page = letter.get_body(("html",)).get_content()
    assert "Подьячий · АО &quot;Ромашка &amp; Ко&quot;" in page and '"Ромашка & Ко"' not in page
    monkeypatch.setattr(mail, "organization", lambda: "")
    assert mail.message(cfg, ["a@example.test"], "тема", "текст", "<p>текст</p>")["Subject"] == "Подьячий: тема"


def test_downloaded_page_gets_the_name_escaped():
    page = mindmap.page("# Протокол\n\n## 1. Кратко\n\nРешили.\n")
    named = mindmap.with_org(page, "Ромашка & <Ко>")
    assert '<span class="meta">Ромашка &amp; &lt;Ко&gt;</span><span class="meta">' in named
    assert named.count("Ромашка") == 1 and len(named) > len(page)


# ---------- сервис ----------

def test_admin_sets_the_logo_and_the_name_and_everyone_sees_them(site):
    import docx
    from pypdf import PdfReader
    web, admin, ivanov, petrov = filled(site)
    guest = site[1]
    assert guest.get("/api/brand").json() == {"name": "", "logo": None}
    seen = admin.get("/api/admin/brand").json()
    assert seen["logo"] == "" and seen["requirements"]["max_bytes"] == 2 * 1024 * 1024
    assert seen["requirements"]["formats"] == ["PNG", "JPEG", "WebP"] and "сценарий" in seen["svg"]

    # обычному пользователю — ни посмотреть настройки, ни поменять
    assert ivanov.get("/api/admin/brand").status_code == 403
    assert ivanov.put("/api/admin/brand/logo", content=picture()).status_code == 403
    assert ivanov.post("/api/admin/brand/name", json={"name": "Иванов и Ко"}).status_code == 403
    assert ivanov.post("/api/admin/brand/reset").status_code == 403
    assert guest.put("/api/admin/brand/logo", content=picture()).status_code == 401

    done = admin.put("/api/admin/brand/logo", content=picture("JPEG", (800, 400), "RGB", (255, 255, 255)))
    assert done.status_code == 200 and len(done.json()["advice"]) == 2 and done.json()["by"] == "admin"
    assert done.json()["source"]["format"] == "JPEG"
    urls = guest.get("/api/brand").json()["logo"]                           # без входа: это видит страница входа
    assert sorted(urls) == ["full", "icon", "small"] and urls == done.json()["public"]["logo"]
    for kind, url in urls.items():
        got = guest.get(url)
        assert got.status_code == 200 and got.headers["content-type"] == "image/png", kind
        assert got.headers["x-content-type-options"] == "nosniff"
        assert got.headers["cache-control"] == "public, max-age=31536000, immutable"
        assert got.content.startswith(b"\x89PNG") and Image.open(io.BytesIO(got.content)).format == "PNG"
    assert Image.open(io.BytesIO(guest.get(urls["icon"]).content)).size == (64, 64)

    # замена логотипа: адрес новый, старый больше не отдаётся — браузеры возьмут новый сразу
    admin.put("/api/admin/brand/logo", content=picture())
    fresh = guest.get("/api/brand").json()["logo"]
    assert fresh["small"] != urls["small"] and guest.get(urls["small"]).status_code == 404
    assert guest.get(fresh["small"]).status_code == 200
    for bad in ("/brand/нет/logo.png", fresh["small"].replace("logo-small.png", "brand.json"),
                fresh["small"].replace("logo-small.png", "..%2Fbrand.json")):
        assert guest.get(bad).status_code == 404, bad

    # подмена расширения и вида: смотрим на содержимое
    for body, why in ((SVG, "может быть сценарий"), (b"GIF89a" + b"0" * 300, "GIF не принимается"),
                      ("просто текст".encode(), "не картинка PNG, JPEG или WebP"), (picture(size=(100, 100)), "слишком мала"),
                      (picture()[:300], "повреждён")):
        refused = admin.put("/api/admin/brand/logo", content=body, headers={"Content-Type": "image/png"})
        assert refused.status_code == 400 and why in refused.json()["detail"], why
    big = admin.put("/api/admin/brand/logo", content=b"\x89PNG\r\n\x1a\n" + b"0" * (2 * 1024 * 1024))
    assert big.status_code == 413 and "больше 2 МБ" in big.json()["detail"]
    assert guest.get("/api/brand").json()["logo"] == fresh                   # отказы ничего не изменили

    for name, why in (("я" * 41, "до 40 знаков"), ("<script>alert(1)</script>", "знаков < и >")):
        refused = admin.post("/api/admin/brand/name", json={"name": name})
        assert refused.status_code == 400 and why in refused.json()["detail"]
    named = admin.post("/api/admin/brand/name", json={"name": '  АО "Ромашка & Ко" '}).json()
    assert named["name"] == 'АО "Ромашка & Ко"' and guest.get("/api/brand").json()["name"] == 'АО "Ромашка & Ко"'

    # название — в колонтитуле скачиваемых файлов; логотип в файлы не встраивается
    query = {"file": "Проекты/вкс.mp4"}
    word = docx.Document(io.BytesIO(ivanov.get("/api/export", params={**query, "format": "docx"}).content))
    assert word.sections[0].header.paragraphs[0].text == 'АО "Ромашка & Ко"' and not word.inline_shapes
    pdf = PdfReader(io.BytesIO(ivanov.get("/api/export", params={**query, "format": "pdf"}).content))
    assert 'АО "Ромашка & Ко"' in pdf.pages[0].extract_text()
    page = ivanov.get("/api/v1/record.html", params=query)
    assert page.status_code == 200 and '<span class="meta">АО "Ромашка &amp; Ко"</span>' in page.text
    assert "attachment" in page.headers["content-disposition"] and "<img" not in page.text.split("<body")[1][:600]
    assert 'Ромашка' in ivanov.get("/api/map", params=query).text
    assert web.core.mail.service_label() == 'Подьячий · АО "Ромашка & Ко"'

    assert [k for k in kinds(web) if k.startswith("brand")] == ["brand_logo", "brand_logo", "brand_name"]
    event = [e for e in web.core.trail.read() if e["kind"] == "brand_logo"][0]
    assert (event["by"], event["format"], event["width"], event["height"]) == ("admin", "JPEG", 800, 400)

    cleared = admin.post("/api/admin/brand/reset").json()
    assert cleared["logo"] == "" and cleared["name"] == "" and guest.get("/api/brand").json() == {"name": "", "logo": None}
    assert guest.get(fresh["small"]).status_code == 404 and web.core.mail.service_label() == "Подьячий"
    word = docx.Document(io.BytesIO(ivanov.get("/api/export", params={**query, "format": "docx"}).content))
    assert word.sections[0].header.paragraphs[0].text == ""
    assert "Ромашка" not in ivanov.get("/api/v1/record.html", params=query).text
    assert [k for k in kinds(web) if k.startswith("brand")][-1] == "brand_reset"
    # в выходной папке и в каталоге приложения ничего не появилось: оформление живёт в папке состояния
    assert not list(web.core.cfg.outbox.rglob("brand*")) and (web.core.cfg.state_dir / "brand").is_dir()


def test_look_changes_only_on_the_page_even_with_a_manage_key(site):
    web, admin, _, _ = staffed(site)
    steward, _ = keyed(web, admin, "с управлением", 30, manage=True)
    robot, _ = keyed(web, admin, "обычный")
    assert steward.get("/api/admin/brand").status_code == 200                    # посмотреть ключом можно
    assert robot.get("/api/admin/brand").json()["detail"] == web.access.BY_KEY_ONLY_READS
    for client, detail in ((steward, web.access.PAGE_ONLY), (robot, web.access.BY_KEY_ONLY_READS)):
        for answer in (client.put("/api/admin/brand/logo", content=picture()),
                       client.post("/api/admin/brand/name", json={"name": "Ромашка"}),
                       client.post("/api/admin/brand/reset")):
            assert answer.status_code == 403 and answer.json()["detail"] == detail
    assert site[1].get("/api/brand").json() == {"name": "", "logo": None} and not (web.core.cfg.state_dir / "brand").exists()


def test_personal_mode_has_no_look_of_its_own(tmp_path):
    with running(tmp_path, mode="personal") as (web, client):
        assert client.get("/api/brand").json() == {"name": "", "logo": None}
        for answer in (client.get("/api/admin/brand"), client.put("/api/admin/brand/logo", content=picture()),
                       client.post("/api/admin/brand/name", json={"name": "Ромашка"}),
                       client.post("/api/admin/brand/reset")):
            assert answer.status_code == 404 and "в личном режиме" in answer.json()["detail"]
        assert client.get("/brand/abc/logo.png").status_code == 404
        # даже если файлы оформления остались от корпоративной установки — личная их не показывает
        web.core.look.set_logo(picture(), "owner")
        web.core.look.set_name("Ромашка", "owner")
        assert client.get("/api/brand").json() == {"name": "", "logo": None} and web.core.organization() == ""
        assert client.get(f"/brand/{web.core.look.view()['logo']}/logo.png").status_code == 404


def test_letters_of_the_service_are_signed_with_the_name(tmp_path, monkeypatch):
    with running(tmp_path, extra=MAIL_ENV) as (web, client):
        sent = []
        monkeypatch.setattr(web.core.mail, "send", lambda cfg, to, letter, *a, **kw: sent.append(
            web.core.mail.message(cfg, to, *letter)) or {})
        admin = enter(web, "admin", ADMIN_PASSWORD)
        admin.post("/api/admin/brand/name", json={"name": "ООО Ромашка"})
        assert admin.post("/api/admin/mail/test").status_code == 200
        assert sent[0]["Subject"].startswith("Подьячий · ООО Ромашка: ")
        assert "Подьячий · ООО Ромашка</div>" in sent[0].get_body(("html",)).get_content()
        assert "<img" not in sent[0].get_body(("html",)).get_content()             # логотипа в письмах нет
