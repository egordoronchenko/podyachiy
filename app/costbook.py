"""Отчёт о расходах файлом: книга Excel и CSV.

Файл — весь отчёт сразу, по периоду и шагу, выбранным на странице: сводка, расходы по времени,
по папкам, людям, записям, этапам и поставщикам, и журнал прогонов. Оба формата собираются из
одних и тех же таблиц (sheets), поэтому не расходятся ни между собой, ни с отчётом на странице:
числа берутся из journal.report, строки журнала — из тех же прогонов.

В книге у каждого разреза свой лист; числа записаны числами, даты — датами. CSV — те же таблицы
подряд в одном файле: строка с названием, шапка, строки, пустая строка между таблицами.

Суммы — в валюте поставщика, по колонке на валюту: рубли и доллары не пересчитываются и
не складываются.
"""
from __future__ import annotations

import csv
import datetime
import io
from dataclasses import dataclass, field

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import journal, usage

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

BUCKETS = {"day": "по дням", "week": "по неделям", "month": "по месяцам", "year": "по годам"}
STAGES = {"speech": "распознавание", "triage": "разбор записи", "actions": "поручения",
          "summary": "протокол", journal.EARLIER: "до появления журнала"}
REASONS = {"new": "новая запись", "retry": "повтор после ошибки", "manual": "ручной запуск",
           "rebuild": "пересборка протокола", journal.EARLIER: "до появления журнала"}
STATUSES = {"ok": "готово", "no_audio": "нет звука", "no_summary": "нет протокола", "error": "ошибка"}
UNKNOWN = "неизвестно"
OUTSIDE = "(вне папок)"

# Вид чисел и дат в книге. Сумма — не меньше двух знаков после запятой и до шести: прогон стоит
# и доли цента, и округлять их в файле для бухгалтерии нельзя.
MONEY, MINUTES, WHOLE = "#,##0.00####", "#,##0.0", "#,##0"
DAY, MOMENT, MONTH = "DD.MM.YYYY", "DD.MM.YYYY HH:MM:SS", "MM.YYYY"
WIDTHS = (8, 70)            # ширина столбца — по содержимому, но в этих пределах


class Month(datetime.date):
    """Месяц целиком: в книге — дата первого числа, показанная как «10.2026»."""


@dataclass
class Table:
    title: str
    head: list[str]
    rows: list[list] = field(default_factory=list)
    money: set[int] = field(default_factory=set)        # номера столбцов с деньгами; прочие дробные — минуты


def name(start: str, end: str, ext: str) -> str:
    return "costs" + "".join(f"_{day}" for day in (start, end) if day) + "." + ext


def _minutes(sec) -> float:
    return round(float(sec or 0) / 60, 1)


def _period(value: str, bucket: str):
    if bucket == "year":
        return int(value)
    if bucket == "month":
        return Month(int(value[:4]), int(value[5:7]), 1)
    return datetime.date.fromisoformat(value)


def sheets(report: dict, runs: list[dict], names: dict[str, str], made_at: datetime.datetime,
           by: str = "") -> list[tuple[str, list[Table]]]:
    """Листы книги: название и таблицы на нём. report — journal.report, runs — прогоны того же
    периода, names — ФИО по логинам, by — кто формирует отчёт."""
    curs = report["currencies"]
    spent = [f"Потрачено, {usage.SIGNS.get(cur, cur)}" for cur in curs]
    money = lambda first: set(range(first, first + len(curs)))
    cost = lambda sums: [float(sums.get(cur, 0.0)) for cur in curs]
    volume = lambda row: [row["records"], _minutes(row["audio_sec"]), row["runs"]]
    person = lambda login: names.get(login, login) if login else UNKNOWN
    total = report["total"]

    day = lambda value, empty: datetime.date.fromisoformat(value) if value else empty
    params = Table("Параметры отчёта", ["Показатель", "Значение"], [
        ["Период с", day(report["start"], "с начала журнала")],
        ["Период по", day(report["end"], "без границы")],
        ["Шаг по времени", BUCKETS[report["bucket"]]],
        ["Отборы", "нет: все папки, люди и поставщики"],
        ["Сформирован", made_at],
        *([["Кем сформирован", f"{names[by]} ({by})" if names.get(by) else by]] if by else []),
        ["Валюты", "суммы — в валюте поставщика; рубли и доллары не пересчитываются и не складываются"],
    ])
    totals = Table("Итоги", ["Показатель", "Значение"], [
        ["Записей", total["records"]], ["Звука, мин", _minutes(total["audio_sec"])], ["Прогонов", total["runs"]]])
    paid = report["paid_audio"]["total"]
    by_currency = Table("По валютам", ["Валюта", "Потрачено", "Оплачено звука, мин", "За час звука"], [
        [cur, float(total["cost"][cur]), _minutes(paid.get(cur)),
         round(total["cost"][cur] / (paid[cur] / 3600), 6) if paid.get(cur) else None]
        for cur in curs], {1, 3})

    periods = Table("По периодам", ["Период", *spent, "Записей", "Звука, мин", "Прогонов"],
                    [[_period(s["period"], report["bucket"]), *cost(s["cost"]), *volume(s)] for s in report["series"]],
                    money(1))

    # записи вне папок в отчёте входят только в общий итог — здесь им своя строка, чтобы лист сходился с итогом
    loose = [r for r in report["records"] if "/" not in r["file"]]
    loose_cost: dict[str, float] = {}
    for r in loose:
        loose_cost = usage.plus(loose_cost, r["cost"])
    folders = Table("Папки", ["Папка", *(f"{s} (с подпапками)" for s in spent), *(f"{s} (без подпапок)" for s in spent),
                              "Записей", "Звука, мин", "Прогонов"],
                    [[f["folder"], *cost(f["cost"]), *cost(f["own"]), *volume(f)] for f in report["folders"]],
                    money(1) | money(1 + len(curs)))
    if loose:
        folders.rows.append([OUTSIDE, *cost(loose_cost), *cost(loose_cost), len(loose),
                             round(sum(_minutes(r["audio_sec"]) for r in loose), 1), sum(r["runs"] for r in loose)])

    users = Table("Пользователи", ["Кто потратил", "Логин", *spent, "Записей", "Звука, мин", "Прогонов"],
                  [[person(u["login"]), u["login"], *cost(u["cost"]), *volume(u)] for u in report["users"]], money(2))
    records = Table("Записи", ["Папка", "Запись", *spent, "Звука, мин", "Прогонов"],
                    [[r["file"].rpartition("/")[0], r["file"].rpartition("/")[2], *cost(r["cost"]),
                      _minutes(r["audio_sec"]), r["runs"]] for r in report["records"]], money(2))
    stages = Table("Этапы и модели", ["Этап", "Модель", *spent, "Прогонов"],
                   [[STAGES.get(s["stage"], s["stage"]), s["model"], *cost(s["cost"]), s["runs"]]
                    for s in report["stages"]], money(2))
    providers = Table("Поставщики", ["Поставщик", *spent, "Записей", "Звука, мин", "Прогонов"],
                      [[p["provider"] or "не назван", *cost(p["cost"]), *volume(p)] for p in report["providers"]],
                      money(1))

    # Строка — то, что прогон заплатил одному поставщику: начатый у основного и законченный у запасного
    # даёт две строки, каждая в своей валюте. Прогон без расхода — одна строка с нулём.
    log = Table("Прогоны", ["Время", "Папка", "Запись", "Кто потратил", "Загрузил", "Запустил", "Причина",
                            "Статус", "Звука, мин", "Поставщик", "Валюта", "Потрачено"], [], {11})
    for run in runs:
        folder, _, record = run["file"].rpartition("/")
        for part in run["providers"] or [{"provider": "", "currency": "", "cost": 0.0}]:
            log.rows.append([datetime.datetime.fromisoformat(run["at"]), folder, record, journal.spender(run),
                             run.get("uploaded_by") or "", run.get("by") or "",
                             REASONS.get(run.get("reason"), run.get("reason") or ""),
                             STATUSES.get(run.get("status"), run.get("status") or ""),
                             _minutes(run.get("audio_sec")), part["provider"], part["currency"], float(part["cost"])])

    return [("Сводка", [params, totals, by_currency]),
            *((t.title, [t]) for t in (periods, folders, users, records, stages, providers, log))]


# ---------- CSV ----------

# Ячейку CSV, которая начинается с одного из этих знаков, Excel и LibreOffice читают как формулу:
# имя записи или папки «=HYPERLINK(…)» стало бы ссылкой наружу, «=cmd|…» — запуском программы.
# Знаки табуляции и перевода строки впереди тоже: программа их отбрасывает и видит формулу за ними.
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def cell(text: str) -> str:
    """Строка для ячейки CSV — всех CSV сервиса: начинается как формула — спереди апостроф, и
    программа покажет её текстом. Числа сюда не идут: они пишутся числами и остаются ими."""
    return "'" + text if text.startswith(FORMULA_START) else text


def _plain(value, money: bool) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.6f}" if money else f"{value:.1f}"
    if isinstance(value, Month):
        return value.isoformat()[:7]
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.isoformat()
    return cell(value) if isinstance(value, str) else str(value)


def to_csv(book: list[tuple[str, list[Table]]]) -> str:
    """Все таблицы подряд: название, шапка, строки, пустая строка. Разделитель «;», дробная часть
    через точку, даты — ГГГГ-ММ-ДД; в начале — метка порядка байтов, чтобы Excel прочитал кириллицу."""
    out = io.StringIO()
    sheet = csv.writer(out, delimiter=";", lineterminator="\r\n")
    for _, tables in book:
        for table in tables:
            sheet.writerow([table.title])
            sheet.writerow(table.head)
            for row in table.rows:
                sheet.writerow([_plain(v, n in table.money) for n, v in enumerate(row)])
            sheet.writerow([])
    return "﻿" + out.getvalue()


# ---------- книга Excel ----------

HEAD_FONT, TITLE_FONT, TOP_FONT = Font(bold=True), Font(bold=True, size=12), Font(bold=True, size=14)
HEAD_FILL = PatternFill("solid", start_color="DCE6F1", end_color="DCE6F1")
HEAD_EDGE = Border(bottom=Side(style="thin", color="8EA9C9"))


def _shown(value, money: bool) -> int:
    """Сколько знаков займёт значение в ячейке — для ширины столбца."""
    if isinstance(value, datetime.datetime):
        return len(MOMENT)
    if isinstance(value, datetime.date):
        return len(DAY)
    if isinstance(value, float):
        return len(f"{value:,.6f}".rstrip("0")) + 2 if money else len(f"{value:,.1f}")
    return len(str(value if value is not None else ""))


def _put(ws, row: int, col: int, value, money: bool) -> None:
    if value is None or value == "":
        return
    if isinstance(value, datetime.datetime):
        value = value.replace(tzinfo=None)          # Excel не знает часовых поясов: время — как записано в журнале
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str):
        cell.data_type = "s"                        # имя записи «=…» — текст, а не формула
    elif isinstance(value, int):
        cell.number_format = WHOLE
    elif isinstance(value, float):
        cell.number_format = MONEY if money else MINUTES
    elif isinstance(value, datetime.datetime):
        cell.number_format = MOMENT
    elif isinstance(value, Month):
        cell.number_format = MONTH
    else:
        cell.number_format = DAY
    if not isinstance(value, (str, int, float)):
        cell.alignment = Alignment(horizontal="left")


def to_xlsx(book: list[tuple[str, list[Table]]]) -> bytes:
    """Книга: по листу на разрез. На листе с одной таблицей шапка закреплена и стоит автофильтр;
    на «Сводке» таблиц несколько — у каждой свой заголовок."""
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = wb.properties.title = "Подьячий — отчёт о расходах"
    for title, tables in book:
        ws = wb.create_sheet(title)
        widths: dict[int, int] = {}
        single = len(tables) == 1
        at = 1
        if not single:
            ws.cell(row=1, column=1, value="Отчёт о расходах").font = TOP_FONT
            at = 3
        for table in tables:
            if not single:
                ws.cell(row=at, column=1, value=table.title).font = TITLE_FONT
                at += 1
            for col, text in enumerate(table.head, 1):
                cell = ws.cell(row=at, column=col, value=text)
                cell.font, cell.fill, cell.border = HEAD_FONT, HEAD_FILL, HEAD_EDGE
                cell.alignment = Alignment(vertical="center", wrap_text=False)
                widths[col] = max(widths.get(col, 0), len(text) + (3 if single else 0))     # место под кнопку фильтра
            head_at = at
            for row in table.rows:
                at += 1
                for col, value in enumerate(row, 1):
                    _put(ws, at, col, value, col - 1 in table.money)
                    widths[col] = max(widths.get(col, 0), _shown(value, col - 1 in table.money))
            if single:
                ws.freeze_panes = ws.cell(row=head_at + 1, column=1)
                ws.auto_filter.ref = f"A{head_at}:{get_column_letter(len(table.head))}{at}"
            at += 2
        for col, width in widths.items():
            ws.column_dimensions[get_column_letter(col)].width = min(max(width + 2, WIDTHS[0]), WIDTHS[1])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
