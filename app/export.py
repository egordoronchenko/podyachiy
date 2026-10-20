"""Протокол записи файлом для людей: DOCX и PDF.

Протокол хранится в Markdown. Здесь он один раз разбирается в блоки (заголовок, абзац,
список, таблица, цитата, код, черта), и из одних и тех же блоков собираются оба файла —
поэтому DOCX и PDF не расходятся ни по содержанию, ни по тому, какая разметка понята.
Разметка — только та, что встречается в протоколах сервиса (см. mindmap.render и mdToHtml
на странице).

Файл строится в памяти и на диск не пишется. Шрифт PDF — DejaVu Sans из app/static/fonts:
в нём есть кириллица, стрелки и типографские знаки; системных шрифтов в образе нет.

Метки времени [чч:мм:сс] и ссылки на другие записи становятся ссылками на страницу сервиса,
только когда известен его постоянный адрес (PUBLIC_URL): файл пересылают и открывают
где угодно, и адрес, взятый из случайного запроса, в нём был бы ложным обещанием. Без
постоянного адреса они остаются обычным текстом.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path, PurePosixPath
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import CondPageBreak, HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from . import external

FORMATS = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}
AUTHOR = "Подьячий"

# ---------- блоки ----------


@dataclass(frozen=True)
class Span:
    """Кусок строки с одним оформлением. href — уже готовый адрес ссылки либо пусто."""
    text: str
    bold: bool = False
    italic: bool = False
    code: bool = False
    href: str = ""
    clock: bool = False        # метка времени [чч:мм:сс]


Line = list[Span]


@dataclass
class Heading:
    level: int
    spans: Line


@dataclass
class Para:
    lines: list[Line]


@dataclass
class Item:
    spans: Line
    level: int = 0             # вложенность: 0 — верхний уровень
    ordered: bool = False
    number: int = 0            # номер пункта, как он стоит в протоколе


@dataclass
class ListBlock:
    items: list[Item]


@dataclass
class Row:
    cells: list[Line]
    group: bool = False        # строка во всю ширину: имя ответственного над его задачами


@dataclass
class TableBlock:
    header: list[Line]
    rows: list[Row]


@dataclass
class Quote:
    lines: list[Line]


@dataclass
class Code:
    text: str


@dataclass
class Rule:
    pass


@dataclass
class Doc:
    title: str
    facts: list[tuple[str, str]] = field(default_factory=list)     # строка сведений под названием
    blocks: list = field(default_factory=list)
    org: str = ""              # название организации из раздела «Оформление» — в верхний колонтитул


# ---------- разбор ----------

TOKEN = re.compile(
    r"\[(?P<text>[^\]\[]+)\]\((?P<url>(?:https?://|/)[^)\s]*)\)"      # [текст](адрес)
    r"|`(?P<code>[^`]+)`"
    r"|\*\*(?P<bold>.+?)\*\*"
    r"|(?<![\w*])\*(?P<italic>[^*\s](?:[^*]*[^*\s])?)\*(?![\w*])"
    r"|\[(?P<clock>\d{1,2}:\d\d:\d\d)\]")                              # метка времени
ITEM = re.compile(r"(\s*)([-*+]|\d+[.)])\s+(.*)")
RULE = re.compile(r"\s*(?:---+|___+|\*\*\*+)\s*")
LEVELS = 3           # глубже третьего уровня списки в протоколах не уходят; лишнее прижимается к третьему


class _Links:
    """Куда ведут ссылки протокола записи file, если у сервиса есть постоянный адрес base."""

    def __init__(self, base: str, file: str):
        self.base, self.file = base, file

    def moment(self, clock: str) -> str:
        return external.page_url(self.file, self.base, external.seconds(clock)) if self.base and self.file else ""

    def page(self, url: str) -> str:
        if url.startswith("/"):
            return self.base + url if self.base else ""
        return url


def _inline(text: str, links: _Links, bold: bool = False, italic: bool = False) -> Line:
    out: Line = []

    def plain(piece: str) -> None:
        if piece:
            out.append(Span(piece, bold, italic))

    at = 0
    for m in TOKEN.finditer(text):
        plain(text[at:m.start()])
        at = m.end()
        if m["url"] is not None:
            out.append(Span(m["text"], bold, italic, href=links.page(m["url"])))
        elif m["code"] is not None:
            out.append(Span(m["code"], bold, italic, code=True))
        elif m["bold"] is not None:
            out.extend(_inline(m["bold"], links, True, italic))
        elif m["italic"] is not None:
            out.extend(_inline(m["italic"], links, bold, True))
        else:
            out.append(Span(m.group(0), bold, italic, href=links.moment(m["clock"]), clock=True))
    plain(text[at:])
    return out


def is_lead(block: Para) -> bool:
    """Абзац-подзаголовок «**Тема — докладчик** [время]»: целиком жирный, под ним идёт его
    список. Заголовком он не размечен, но отрываться от списка на другую страницу не должен."""
    said = [s for line in block.lines for s in line if s.text.strip() and not s.clock]
    return len(block.lines) == 1 and bool(said) and all(s.bold for s in said)


def text_of(line: Line) -> str:
    return "".join(s.text for s in line)


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def _table(block: list[list[str]], links: _Links) -> TableBlock:
    head, body = block[0], [r for r in block[1:] if not all(re.fullmatch(r":?-+:?", c or "-") for c in r)]
    width = max(len(r) for r in [head] + body)
    rows = []
    for r in body:
        group = bool(re.fullmatch(r"\*\*.+\*\*", r[0])) and not any(c for c in r[1:])
        r = r + [""] * (width - len(r))
        rows.append(Row([_inline(c, links) for c in r], group))
    return TableBlock([_inline(c, links) for c in head + [""] * (width - len(head))], rows)


def parse(md: str, base: str = "", file: str = "") -> list:
    """Протокол в блоки. base — постоянный адрес сервиса, file — запись: с ними метки времени
    и ссылки на другие записи получают адреса, без них остаются текстом."""
    links = _Links(base, file)
    lines = re.sub(r"<!--.*?-->", "", md.replace("\r\n", "\n"), flags=re.S).split("\n")
    blocks: list = []
    para: list[Line] = []
    items: list[Item] | None = None
    indents: list[int] = []          # отступы открытых уровней списка

    def close() -> None:
        nonlocal items
        if para:
            blocks.append(Para(list(para)))
            para.clear()
        if items:
            blocks.append(ListBlock(items))
        items = None
        indents.clear()

    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if ln.lstrip().startswith("```"):
            close()
            code = []
            i += 1
            while i < len(lines) and not lines[i].lstrip().startswith("```"):
                code.append(lines[i].rstrip())
                i += 1
            blocks.append(Code("\n".join(code)))
        elif m := re.match(r"(#{1,6})\s+(.*)", ln):
            close()
            blocks.append(Heading(min(len(m.group(1)), 4), _inline(m.group(2), links)))
        elif ln.lstrip().startswith("|"):
            close()
            block = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                block.append(_cells(lines[i]))
                i += 1
            i -= 1
            blocks.append(_table(block, links))
        elif RULE.fullmatch(ln):
            close()
            blocks.append(Rule())
        elif ln.lstrip().startswith(">"):
            close()
            quote = []
            while i < len(lines) and lines[i].lstrip().startswith(">"):
                quote.append(_inline(re.sub(r"^\s*>\s?", "", lines[i].rstrip()), links))
                i += 1
            i -= 1
            blocks.append(Quote(quote))
        elif m := ITEM.match(ln):
            if para:
                close()
            indent, mark = len(m.group(1)), m.group(2)
            while indents and indents[-1] > indent:
                indents.pop()
            if not indents or indents[-1] < indent:
                indents.append(indent)
            if items is None:
                items = []
            ordered = mark[0].isdigit()
            items.append(Item(_inline(m.group(3), links), min(len(indents), LEVELS) - 1, ordered,
                              int(mark[:-1]) if ordered else 0))
        elif not ln.strip():
            nxt = lines[i + 1] if i + 1 < len(lines) else ""
            if not (items and (ITEM.match(nxt) or nxt.startswith(" "))):
                close()
            elif para:
                close()
        elif items and ln.startswith(" "):
            last = items[-1]                          # продолжение пункта на следующей строке
            last.spans = last.spans + _inline(" " + ln.strip(), links)
        else:
            if items:
                close()
            para.append(_inline(ln.strip(), links))
        i += 1
    close()
    return blocks


def clock(sec: float) -> str:
    """Длительность словами: «1 ч 27 мин», «31 мин»."""
    minutes = max(1, round(sec / 60))
    return f"{minutes // 60} ч {minutes % 60:02d} мин" if minutes >= 60 else f"{minutes} мин"


def document(md: str, record: dict, base: str = "", org: str = "") -> Doc:
    """Документ для выгрузки: название записи, строка сведений о ней и блоки протокола."""
    file = record["file"]
    path = PurePosixPath(file)
    blocks = parse(md, base, file)
    first = next((b for b in blocks if isinstance(b, Heading) and b.level == 1), None)
    title = (record.get("title") or "").strip() or (text_of(first.spans).strip() if first else "") or path.stem
    if first is not None and text_of(first.spans).strip() == title:
        blocks.remove(first)                          # название уже стоит заголовком документа
    facts = []
    day = external.day_of(record)
    if day:
        named = re.fullmatch(r"\d{4}-\d\d-\d\d", record.get("record_date") or "")
        facts.append(("Дата совещания" if named else "Обработано", f"{day[8:10]}.{day[5:7]}.{day[:4]}"))
    if record.get("audio_duration_sec"):
        facts.append(("Длительность", clock(record["audio_duration_sec"])))
    if path.parent.as_posix() != ".":
        facts.append(("Папка", path.parent.as_posix()))
    facts.append(("Запись", path.name))
    return Doc(title, facts, blocks, org)


def file_name(title: str, fmt: str) -> str:
    """Имя скачиваемого файла — название записи без знаков, которых не терпят файловые системы."""
    name = re.sub(r'[\\/:*?"<>|\x00-\x1f]+', " ", title)
    name = re.sub(r"\s+", " ", name).strip(" .")[:120].rstrip(" .")
    return f"{name or 'протокол'}.{fmt}"


def shares(table: TableBlock) -> list[float]:
    """Доли ширины колонок. Колонка получает место по самому длинному тексту в ней, но
    длинные тексты растут медленнее коротких: «№» и метка времени остаются узкими и при
    этом не ломаются посреди слова, а описание задачи забирает остальное."""
    weights = []
    for i in range(len(table.header)):
        texts = [text_of(table.header[i])] + [text_of(r.cells[i]) for r in table.rows if not r.group]
        longest = max(len(t) for t in texts)
        word = max((len(w) for t in texts for w in t.split()), default=1)
        weights.append(max(min(word, 14) + 3.0, min(longest, 70) ** 0.8, 5.0))
    total = sum(weights)
    return [w / total for w in weights]


# ---------- общее оформление ----------

PAGE_W, PAGE_H = 21.0, 29.7                    # А4, см
LEFT, RIGHT, TOP, BOTTOM = 2.5, 1.5, 2.0, 2.0  # поля, см: слева шире — под подшивку
INK, DIM, LINK_COLOR = "1F2937", "5B6472", "1D4ED8"
HEAD_FILL, GROUP_FILL, GRID = "E8ECF2", "F3F5F8", "A9B1BD"
BULLETS = ("•", "–", "◦")


# ---------- DOCX ----------

FACE = "Arial"                 # есть в Windows; LibreOffice и «МойОфис» подставляют совместимый по ширине
MONO = "Courier New"
HEADING_PT = {1: 15, 2: 13, 3: 11.5, 4: 10.5}
LIST_STYLES = {False: ("List Bullet", "List Bullet 2", "List Bullet 3"),
               True: ("List Number", "List Number 2", "List Number 3")}


def _face(style, name: str, size: float, color: str = INK, bold: bool | None = None) -> None:
    """Шрифт стиля — явно для всех письменностей: иначе Word берёт шрифт темы, а не этот."""
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        style.font.bold = bold
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in list(fonts.attrib):
        if "Theme" in key:
            del fonts.attrib[key]
    for kind in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(kind), name)


def _docx_styles(doc) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    _face(normal, FACE, 10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    lang = OxmlElement("w:lang")                       # язык проверки орфографии
    lang.set(qn("w:val"), "ru-RU")
    normal.element.get_or_add_rPr().append(lang)

    title = styles["Title"]
    _face(title, FACE, 18, bold=True)
    title.paragraph_format.space_after = Pt(4)
    border = title.element.pPr.find(qn("w:pBdr"))      # у стиля шаблона под названием синяя черта
    if border is not None:
        title.element.pPr.remove(border)
    for level, size in HEADING_PT.items():
        head = styles[f"Heading {level}"]
        _face(head, FACE, size, bold=True)
        head.font.italic = False
        head.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
        head.paragraph_format.space_after = Pt(5)
        head.paragraph_format.keep_with_next = True
    for names in LIST_STYLES.values():
        for name in names:
            _face(styles[name], FACE, 10.5)
            styles[name].paragraph_format.space_after = Pt(3)
    _face(styles["Quote"], FACE, 10.5, DIM)
    styles["Quote"].paragraph_format.left_indent = Cm(0.8)
    _face(styles["Footer"], FACE, 8.5, DIM)
    _face(styles["Header"], FACE, 8.5, DIM)

    facts = styles.add_style("Protocol Facts", 1)
    facts.base_style = normal
    _face(facts, FACE, 9, DIM)
    facts.paragraph_format.space_after = Pt(10)
    cell = styles.add_style("Protocol Cell", 1)
    cell.base_style = normal
    _face(cell, FACE, 9)
    cell.paragraph_format.space_before = Pt(2)
    cell.paragraph_format.space_after = Pt(2)
    cell.paragraph_format.line_spacing = 1.0
    code = styles.add_style("Protocol Code", 1)
    code.base_style = normal
    _face(code, MONO, 9)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.left_indent = Cm(0.5)


def _docx_run(par, span: Span, bold: bool = False):
    run = par.add_run(span.text)
    run.bold = True if (span.bold or bold) else None
    run.italic = True if span.italic else None
    if span.code:
        run.font.name = MONO
        fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
        for kind in ("w:cs", "w:eastAsia"):
            fonts.set(qn(kind), MONO)
    if span.href:
        run.font.color.rgb = RGBColor.from_string(LINK_COLOR)
        run.font.underline = not span.clock        # меток времени в протоколе сотни: подчёркивание рябит
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), par.part.relate_to(span.href, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
        par._p.append(link)
        link.append(run._r)
    return run


def _docx_line(par, line: Line, bold: bool = False) -> None:
    for span in line:
        _docx_run(par, span, bold)


def _docx_lines(par, lines: list[Line]) -> None:
    for n, line in enumerate(lines):
        if n:
            par.add_run().add_break()
        _docx_line(par, line)


def _field(par, code: str) -> None:
    """Поле Word: номер страницы, число страниц."""
    el = OxmlElement("w:fldSimple")
    el.set(qn("w:instr"), code)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    el.append(run)
    par._p.append(el)


def _shade(cell, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _docx_table(doc, block: TableBlock, width_cm: float) -> None:
    cols = len(block.header)
    table = doc.add_table(rows=1, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    borders = OxmlElement("w:tblBorders")               # тонкая серая сетка вместо чёрной из стиля
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{side}")
        for key, value in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", GRID)):
            edge.set(qn(f"w:{key}"), value)
        borders.append(edge)
    table._tbl.tblPr.append(borders)
    widths = [Cm(width_cm * s) for s in shares(block)]
    for column, width in zip(table.columns, widths):
        column.width = width

    def fill(cells, lines: list[Line], bold: bool, shade: str = "") -> None:
        for cell, width, line in zip(cells, widths, lines):
            cell.width = width
            par = cell.paragraphs[0]
            par.style = "Protocol Cell"
            _docx_line(par, line, bold)
            if shade:
                _shade(cell, shade)

    head = table.rows[0]
    fill(head.cells, block.header, True, HEAD_FILL)
    pr = head._tr.get_or_add_trPr()
    pr.append(OxmlElement("w:cantSplit"))
    pr.append(OxmlElement("w:tblHeader"))               # шапка повторяется на каждой странице
    for row in block.rows:
        added = table.add_row()
        added._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))     # строка не рвётся между страницами
        cells = added.cells
        if row.group:
            merged = cells[0].merge(cells[-1])
            par = merged.paragraphs[0]
            par.style = "Protocol Cell"
            par.paragraph_format.keep_with_next = True      # имя не остаётся внизу страницы без своих задач
            _docx_line(par, row.cells[0], True)
            _shade(merged, GROUP_FILL)
        else:
            fill(cells, row.cells, False)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


class _Numbering:
    """Нумерованные списки Word. У стиля «List Number» один счётчик на весь документ, и второй
    список продолжил бы первый — поэтому каждому списку заводится свой счётчик, начинающийся
    с того номера, который стоит в протоколе."""

    def __init__(self, doc):
        self.doc = doc
        self.root = doc.part.numbering_part.element
        self.open: dict[int, int] = {}                  # уровень → счётчик идущего списка

    def reset(self) -> None:
        self.open.clear()

    def apply(self, par, item: Item) -> None:
        for level in [k for k in self.open if k > item.level]:
            del self.open[level]
        if not item.ordered:
            self.open.pop(item.level, None)
            return
        if item.level not in self.open:
            style = self.doc.styles[LIST_STYLES[True][item.level]]
            source = self.root.num_having_numId(style.element.pPr.numPr.numId.val)
            num = self.root.add_num(source.abstractNumId.val)
            num.add_lvlOverride(ilvl=0).add_startOverride(item.number)
            self.open[item.level] = num.numId
        pr = par._p.get_or_add_pPr().get_or_add_numPr()
        pr.get_or_add_ilvl().val = 0
        pr.get_or_add_numId().val = self.open[item.level]


def to_docx(doc_: Doc) -> bytes:
    doc = Document()
    _docx_styles(doc)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(PAGE_W), Cm(PAGE_H)
    section.left_margin, section.right_margin = Cm(LEFT), Cm(RIGHT)
    section.top_margin, section.bottom_margin = Cm(TOP), Cm(BOTTOM)
    props = doc.core_properties
    props.title, props.author, props.language = doc_.title, AUTHOR, "ru-RU"
    props.comments = ""

    if doc_.org:
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header.add_run(doc_.org)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Страница ")
    _field(footer, "PAGE")
    footer.add_run(" из ")
    _field(footer, "NUMPAGES")

    doc.add_paragraph(doc_.title, style="Title")
    if doc_.facts:
        doc.add_paragraph(" · ".join(f"{k}: {v}" for k, v in doc_.facts), style="Protocol Facts")

    numbering = _Numbering(doc)
    width = PAGE_W - LEFT - RIGHT
    for block in doc_.blocks:
        if isinstance(block, Heading):
            _docx_line(doc.add_paragraph(style=f"Heading {block.level}"), block.spans)
        elif isinstance(block, Para):
            par = doc.add_paragraph()
            par.paragraph_format.keep_with_next = is_lead(block)
            _docx_lines(par, block.lines)
        elif isinstance(block, ListBlock):
            numbering.reset()
            for item in block.items:
                par = doc.add_paragraph(style=LIST_STYLES[item.ordered][item.level])
                numbering.apply(par, item)
                _docx_line(par, item.spans)
        elif isinstance(block, TableBlock):
            _docx_table(doc, block, width)
        elif isinstance(block, Quote):
            _docx_lines(doc.add_paragraph(style="Quote"), block.lines)
        elif isinstance(block, Code):
            for n, line in enumerate(block.text.split("\n")):
                par = doc.add_paragraph(style="Protocol Code")
                par.add_run(line)
            par.paragraph_format.space_after = Pt(6)
        elif isinstance(block, Rule):
            par = doc.add_paragraph()
            borders = OxmlElement("w:pBdr")
            line = OxmlElement("w:bottom")
            for key, value in (("val", "single"), ("sz", "4"), ("space", "1"), ("color", GRID)):
                line.set(qn(f"w:{key}"), value)
            borders.append(line)
            par._p.get_or_add_pPr().append(borders)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ---------- PDF ----------

FONTS = Path(__file__).parent / "static" / "fonts"
PDF_FACE, PDF_MONO = "DejaVuSans", "DejaVuSansMono"
for _name, _file in ((PDF_FACE, "DejaVuSans.ttf"), (f"{PDF_FACE}-Bold", "DejaVuSans-Bold.ttf"),
                     (f"{PDF_FACE}-Oblique", "DejaVuSans-Oblique.ttf"),
                     (f"{PDF_FACE}-BoldOblique", "DejaVuSans-BoldOblique.ttf"),
                     (PDF_MONO, "DejaVuSansMono.ttf")):
    pdfmetrics.registerFont(TTFont(_name, str(FONTS / _file)))
pdfmetrics.registerFontFamily(PDF_FACE, normal=PDF_FACE, bold=f"{PDF_FACE}-Bold",
                              italic=f"{PDF_FACE}-Oblique", boldItalic=f"{PDF_FACE}-BoldOblique")


def _hex(value: str):
    return colors.HexColor("#" + value)


BODY = ParagraphStyle("body", fontName=PDF_FACE, fontSize=9.5, leading=13.5, textColor=_hex(INK),
                      alignment=TA_LEFT, spaceAfter=6)
PDF_LEAD = ParagraphStyle("lead", parent=BODY, keepWithNext=1, spaceBefore=4)
PDF_TITLE = ParagraphStyle("title", parent=BODY, fontName=f"{PDF_FACE}-Bold", fontSize=17, leading=21, spaceAfter=5)
PDF_FACTS = ParagraphStyle("facts", parent=BODY, fontSize=8.5, leading=12, textColor=_hex(DIM), spaceAfter=12)
PDF_HEAD = {level: ParagraphStyle(f"h{level}", parent=BODY, fontName=f"{PDF_FACE}-Bold", fontSize=size,
                                  leading=size * 1.25, spaceBefore=14 if level <= 2 else 10, spaceAfter=5,
                                  keepWithNext=1)
            for level, size in ((1, 14), (2, 12), (3, 10.5), (4, 9.5))}
# Заголовок прямо перед таблицей: «не отрывать от следующего» здесь значило бы переносить на новую
# страницу всю таблицу целиком, оставляя полстраницы пустой. Вместо этого заголовок просит
# места на себя и первые строки таблицы (TABLE_START), а таблица дальше делится как обычно.
PDF_HEAD_TABLE = {level: ParagraphStyle(f"h{level}t", parent=style, keepWithNext=0)
                  for level, style in PDF_HEAD.items()}
TABLE_START = 4 * cm
PDF_ITEM = [ParagraphStyle(f"li{level}", parent=BODY, leftIndent=16 + 16 * level, bulletIndent=3 + 16 * level,
                           bulletFontName=PDF_FACE, bulletFontSize=9.5, spaceAfter=3)
            for level in range(LEVELS)]
PDF_CELL = ParagraphStyle("cell", parent=BODY, fontSize=8, leading=10.5, spaceAfter=0)
PDF_CELL_BOLD = ParagraphStyle("cellb", parent=PDF_CELL, fontName=f"{PDF_FACE}-Bold")
PDF_QUOTE = ParagraphStyle("quote", parent=BODY, leftIndent=18, textColor=_hex(DIM))
PDF_CODE = ParagraphStyle("code", parent=BODY, fontName=PDF_MONO, fontSize=8, leading=11, leftIndent=12)
CELL_PAD = 4
FRAME_H = (PAGE_H - TOP - BOTTOM) * cm - 12        # высота пустой страницы: поля и отступы рамки по 6 пт


def _markup(line: Line) -> str:
    """Строка в разметке абзаца reportlab."""
    out = []
    for span in line:
        text = escape(span.text)
        if span.code:
            text = f'<font face="{PDF_MONO}">{text}</font>'
        if span.bold:
            text = f"<b>{text}</b>"
        if span.italic:
            text = f"<i>{text}</i>"
        if span.href:
            text = text if span.clock else f"<u>{text}</u>"
            text = f'<a href="{escape(span.href, {chr(34): "&quot;"})}" color="#{LINK_COLOR}">{text}</a>'
        out.append(text)
    return "".join(out)


def _numbered(title: str, org: str = ""):
    """Холст, который знает число страниц: колонтитул «Страница N из M» рисуется в конце,
    когда документ уже разложен."""

    class Numbered(canvas.Canvas):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self.pages: list[dict] = []

        def showPage(self):
            self.pages.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            for page in self.pages:
                self.__dict__.update(page)
                self.setFont(PDF_FACE, 7.5)
                self.setFillColor(_hex(DIM))
                room = (PAGE_W - LEFT - RIGHT) * cm - 90
                name = title
                while name and self.stringWidth(name, PDF_FACE, 7.5) > room:
                    name = name[:-2].rstrip() + "…"
                self.drawString(LEFT * cm, 1.1 * cm, name)
                if org:
                    self.drawRightString((PAGE_W - RIGHT) * cm, (PAGE_H - 1.1) * cm, org)
                self.drawRightString((PAGE_W - RIGHT) * cm, 1.1 * cm,
                                     f"Страница {self._pageNumber} из {len(self.pages)}")
                super().showPage()
            super().save()

    return Numbered


class _Pdf(SimpleDocTemplate):
    """Документ с закладками: заголовки протокола попадают в оглавление просмотрщика."""

    def beforeDocument(self):
        self.marks, self.depth = 0, -1

    def afterFlowable(self, flowable):
        level = getattr(flowable, "outline_level", None)
        if level is None:
            return
        level = min(level, self.depth + 1)              # оглавление PDF не терпит пропуска уровня
        self.depth = level
        self.marks += 1
        key = f"mark{self.marks}"
        self.canv.bookmarkPage(key)
        self.canv.addOutlineEntry(flowable.getPlainText(), key, level, closed=level >= 2)


class _Table(Table):
    """Таблица, которая переносится на следующую страницу целыми строками. Строку режет
    посередине только тогда, когда она не помещается и на пустую страницу: иначе reportlab
    делит строку у нижнего края каждой страницы и ставит перед обрезком вторую шапку."""

    def split(self, availWidth, availHeight):
        self.splitInRow = 0
        parts = super().split(availWidth, availHeight)
        if parts or availHeight < FRAME_H:
            return parts
        self.splitInRow = 1
        return super().split(availWidth, availHeight)


def _pdf_table(block: TableBlock, width: float) -> Table:
    data = [[Paragraph(_markup(c), PDF_CELL_BOLD) for c in block.header]]
    style = [("GRID", (0, 0), (-1, -1), 0.5, _hex(GRID)),
             ("BACKGROUND", (0, 0), (-1, 0), _hex(HEAD_FILL)),
             ("VALIGN", (0, 0), (-1, -1), "TOP"),
             ("LEFTPADDING", (0, 0), (-1, -1), CELL_PAD), ("RIGHTPADDING", (0, 0), (-1, -1), CELL_PAD),
             ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]
    for n, row in enumerate(block.rows, 1):
        if row.group:
            data.append([Paragraph(_markup(row.cells[0]), PDF_CELL_BOLD)] + [""] * (len(block.header) - 1))
            style += [("SPAN", (0, n), (-1, n)), ("BACKGROUND", (0, n), (-1, n), _hex(GROUP_FILL))]
            if n < len(block.rows):                     # имя не остаётся внизу страницы без своих задач
                style.append(("NOSPLIT", (0, n), (-1, n + 1)))
        else:
            data.append([Paragraph(_markup(c), PDF_CELL) for c in row.cells])
    # шапка повторяется на каждой странице
    table = _Table(data, colWidths=[width * s for s in shares(block)], repeatRows=1,
                   hAlign="LEFT", spaceAfter=10)
    table.setStyle(TableStyle(style))
    return table


def to_pdf(doc_: Doc) -> bytes:
    out = io.BytesIO()
    pdf = _Pdf(out, pagesize=A4, leftMargin=LEFT * cm, rightMargin=RIGHT * cm, topMargin=TOP * cm,
               bottomMargin=BOTTOM * cm, title=doc_.title, author=AUTHOR, creator=AUTHOR, lang="ru-RU")
    title = Paragraph(escape(doc_.title), PDF_TITLE)
    title.outline_level = 0
    story: list = [title]
    if doc_.facts:
        story.append(Paragraph(escape(" · ".join(f"{k}: {v}" for k, v in doc_.facts)), PDF_FACTS))
    for block, after in zip(doc_.blocks, doc_.blocks[1:] + [None]):
        if isinstance(block, Heading):
            before_table = isinstance(after, TableBlock)
            if before_table:
                story.append(CondPageBreak(TABLE_START))
            head = Paragraph(_markup(block.spans), (PDF_HEAD_TABLE if before_table else PDF_HEAD)[block.level])
            head.outline_level = block.level
            story.append(head)
        elif isinstance(block, Para):
            story.append(Paragraph("<br/>".join(_markup(line) for line in block.lines),
                                   PDF_LEAD if is_lead(block) else BODY))
        elif isinstance(block, ListBlock):
            for item in block.items:
                mark = f"{item.number}." if item.ordered else BULLETS[item.level]
                story.append(Paragraph(_markup(item.spans), PDF_ITEM[item.level], bulletText=mark))
            story.append(Spacer(1, 4))
        elif isinstance(block, TableBlock):
            story.append(_pdf_table(block, pdf.width))
        elif isinstance(block, Quote):
            story.append(Paragraph("<br/>".join(_markup(line) for line in block.lines), PDF_QUOTE))
        elif isinstance(block, Code):
            lines = [escape(line).replace(" ", "\u00a0") for line in block.text.split("\n")]
            story.append(Paragraph("<br/>".join(lines), PDF_CODE))
        elif isinstance(block, Rule):
            story.append(HRFlowable(width="100%", thickness=0.5, color=_hex(GRID), spaceBefore=4, spaceAfter=8))
    pdf.build(story, canvasmaker=_numbered(doc_.title, doc_.org))
    return out.getvalue()


BUILDERS = {"docx": to_docx, "pdf": to_pdf}


def build(md: str, record: dict, fmt: str, base: str = "", org: str = "") -> tuple[bytes, str]:
    """Файл протокола и его имя. fmt — ключ FORMATS; base — постоянный адрес сервиса или пусто;
    org — название организации для верхнего колонтитула или пусто."""
    doc = document(md, record, base, org)
    return BUILDERS[fmt](doc), file_name(doc.title, fmt)
